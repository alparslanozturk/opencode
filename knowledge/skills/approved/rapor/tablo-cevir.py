#!/usr/bin/env python3
# tablo-cevir.py — betiklerin ürettiği CSV'yi (filo-tablo --csv, sar-ozet --csv-dosya, kume-bul …) değerlerine DOKUNMADAN
# yazdırılabilir HTML'e ya da (openpyxl kuruluysa) XLSX'e çevirir. Sorunlu satırları vurgular, başlığa kaynak + tarih yazar.
# Neden (saha G22, M-1): model rapor kodunu yazarken değerleri koda elle kopyalıyordu — sunucular arası kopyalanan sayılar,
# MB/GB karışımı, tutmayan toplamlar. Burada hiçbir değer yeniden yazılmaz; hücre = CSV hücresi.
#
# kullanım:
#   tablo-cevir.py <girdi.csv> --html <cikti.html> [--baslik "…"] [--kaynak "…komut / betik…"] [--vurgu <sütun>=<regex>]
#   tablo-cevir.py <girdi.csv> --xlsx <cikti.xlsx> [aynı seçenekler]
#   --vurgu verilmezse: `durum` sütunu varsa "ok" olmayan satırlar, `sapan_alanlar` doluysa o satır vurgulanır.
#   Var olan çıktının üzerine yazmaz (kullanıcıya sor, ada tarih ekle).
import csv
import datetime
import html
import io
import os
import re
import sys


def oku(yol):
    ham = open(yol, encoding="utf-8-sig", errors="replace").read()
    ilk = ham.splitlines()[0] if ham else ""
    ayrac = ";" if ilk.count(";") >= ilk.count(",") else ","
    satirlar = list(csv.reader(io.StringIO(ham), delimiter=ayrac))
    return satirlar[0], [s for s in satirlar[1:] if any(x.strip() for x in s)]


def sayi(v):
    v = v.strip()
    if re.match(r"^-?\d+$", v):
        return int(v)
    if re.match(r"^-?\d+\.\d+$", v):
        return float(v)
    return v


def vurgulu_mu(baslik, satir, vurgu):
    d = dict(zip(baslik, satir))
    if vurgu:
        sut, desen = vurgu
        return bool(re.search(desen, d.get(sut, "")))
    if "durum" in d and d["durum"].strip() and d["durum"].strip().lower() != "ok":
        return True
    if d.get("sapan_alanlar", "").strip():
        return True
    return False


def main(argv):
    if not argv or argv[0] in ("-h", "--help"):
        print(open(__file__, encoding="utf-8").read().split("import csv")[0])
        return 0 if argv else 2
    girdi = argv[0]
    sec = lambda ad: argv[argv.index(ad) + 1] if ad in argv and argv.index(ad) + 1 < len(argv) else None
    hedef_html, hedef_xlsx = sec("--html"), sec("--xlsx")
    baslik_metin = sec("--baslik") or os.path.splitext(os.path.basename(girdi))[0]
    kaynak = sec("--kaynak") or "kaynak belirtilmedi — --kaynak ile komut/betik yaz"
    vurgu = None
    if sec("--vurgu"):
        k, _, d = sec("--vurgu").partition("=")
        vurgu = (k, d)
    if not (hedef_html or hedef_xlsx):
        print("--html <dosya> ya da --xlsx <dosya> ver")
        return 2
    hedef = hedef_html or hedef_xlsx
    if os.path.exists(hedef):
        print("%s zaten var — üzerine yazılmadı (kullanıcıya sor, ada tarih ekle)" % hedef)
        return 1
    baslik, satirlar = oku(girdi)
    vurgular = [vurgulu_mu(baslik, s, vurgu) for s in satirlar]
    zaman = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    if hedef_html:
        govde = []
        for s, v in zip(satirlar, vurgular):
            govde.append('<tr%s>%s</tr>' % (' class="sorun"' if v else "", "".join(
                '<td%s>%s</td>' % (' class="sayi"' if isinstance(sayi(x), (int, float)) else "", html.escape(x)) for x in s)))
        sayfa = """<!doctype html><html lang="tr"><head><meta charset="utf-8"><title>%(b)s</title>
<style>
 body { font-family: sans-serif; margin: 2rem; }
 table { border-collapse: collapse; font-size: 13px; }
 th, td { border: 1px solid #ccc; padding: 4px 8px; }
 th { background: #eee; text-align: left; position: sticky; top: 0; }
 td.sayi { text-align: right; font-variant-numeric: tabular-nums; }
 .sorun { background: #fdd; }
 @media print { .sorun { background: #fdd !important; -webkit-print-color-adjust: exact; print-color-adjust: exact; } }
 .not { color: #555; font-size: 12px; }
</style></head><body>
<h1>%(b)s</h1>
<p class="not">Üretildi: %(z)s · Kaynak: %(k)s · Girdi: %(g)s · %(n)d satır, %(v)d vurgulu (kırmızı)</p>
<table><tr>%(th)s</tr>
%(tr)s
</table></body></html>
""" % {"b": html.escape(baslik_metin), "z": zaman, "k": html.escape(kaynak), "g": html.escape(os.path.basename(girdi)),
       "n": len(satirlar), "v": sum(vurgular), "th": "".join("<th>%s</th>" % html.escape(h) for h in baslik), "tr": "\n".join(govde)}
        open(hedef_html, "w", encoding="utf-8").write(sayfa)
    else:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            print("openpyxl kurulu değil — XLSX üretilemedi. --html kullan (tarayıcıdan yazdır → PDF) ya da CSV'yi Excel'de aç.")
            return 1
        wb = Workbook()
        ws = wb.active
        ws.title = baslik_metin[:31]
        ws.append(baslik)
        for h in ws[1]:
            h.font = Font(bold=True)
        kirmizi = PatternFill("solid", start_color="FFC7CE")
        for s, v in zip(satirlar, vurgular):
            ws.append([sayi(x) for x in s])
            if v:
                for h in ws[ws.max_row]:
                    h.fill = kirmizi
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(baslik, 1):
            gen = max([len(h)] + [len(s[i - 1]) for s in satirlar if len(s) >= i])
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, gen + 2))
        not_ = wb.create_sheet("Kaynak")
        for satir in (["Üretildi", zaman], ["Kaynak", kaynak], ["Girdi", os.path.basename(girdi)],
                      ["Satır", len(satirlar)], ["Vurgulu", sum(vurgular)]):
            not_.append(satir)
        wb.save(hedef_xlsx)
    print("yazıldı %s · %d satır · %d vurgulu · sütunlar: %s" % (os.path.abspath(hedef), len(satirlar), sum(vurgular), ", ".join(baslik)))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
